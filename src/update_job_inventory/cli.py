"""CLI: annotate a job-inventory workbook with Original Path / Updated Path columns.

Reads an .xlsx job inventory, derives each row's original directory from its
Command Line (Task Type "Command") or its File Path/Member Library / Embedded
Script (Task Type "Job"), optionally fills Updated Path, and writes a new workbook.
"""

from __future__ import annotations

import argparse
import logging
import posixpath
import re
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

PROG_NAME = "update-job-inventory"
LOG_FILE = f"{PROG_NAME}.log"

ORIGINAL_PATH = "Original Path"
UPDATED_PATH = "Updated Path"

TASK_TYPE = "Task Type"
LIBRARY = "File Path/Member Library"
SCRIPT = "Embedded Script"
COMMAND = "Command Line"

# An absolute POSIX path (/a/b), home-relative (~/a), or Windows path (C:\a or C:/a).
_PATH_RE = re.compile(r"""(?<![\w./~\\-])(?:~?/|[A-Za-z]:[\\/])[^\s"'`|;&<>()]+""")

logger = logging.getLogger(__name__)


def first_directory(text: object) -> str | None:
    """Return the directory of the first absolute path found in ``text``, or None."""
    if not isinstance(text, str):
        return None
    match = _PATH_RE.search(text)
    if match is None:
        return None
    path = match.group(0).rstrip(".,:")
    if "\\" in path:
        path = path.replace("\\", "/")
    if path.endswith("/") and len(path) > 1:
        return path.rstrip("/")
    return posixpath.dirname(path) or "/"


def _original_path(task_type: object, row: dict[str, object]) -> str | None:
    kind = str(task_type or "").strip().lower()
    if kind == "command":
        return first_directory(row.get(COMMAND))
    if kind == "job":
        return first_directory(row.get(LIBRARY)) or first_directory(row.get(SCRIPT))
    return None


def _column_index(ws: Worksheet, headers: dict[str, int], name: str) -> int:
    """Return the 1-based column for ``name``, appending it on the right if missing."""
    if name not in headers:
        col = ws.max_column + 1
        ws.cell(row=1, column=col, value=name)
        headers[name] = col
    return headers[name]


def update_workbook(ws: Worksheet, updated_path: str | None) -> tuple[int, int]:
    """Fill the path columns in place. Returns (rows processed, rows with an original path)."""
    headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    if TASK_TYPE not in headers:
        raise SystemExit(f"error: workbook is missing required column: {TASK_TYPE}")
    source_cols = {h: headers[h] for h in (LIBRARY, SCRIPT, COMMAND) if h in headers}
    orig_col = _column_index(ws, headers, ORIGINAL_PATH)
    upd_col = _column_index(ws, headers, UPDATED_PATH)

    processed = found = 0
    for r in range(2, ws.max_row + 1):
        values = {h: ws.cell(row=r, column=c).value for h, c in source_cols.items()}
        task_type = ws.cell(row=r, column=headers[TASK_TYPE]).value
        if task_type is None and not any(values.values()):
            continue  # blank row
        processed += 1
        original = _original_path(task_type, values)
        if original:
            found += 1
        else:
            logger.warning("Row %d: no path found (Task Type %r)", r, task_type)
        ws.cell(row=r, column=orig_col, value=original)
        if updated_path is not None:
            ws.cell(row=r, column=upd_col, value=updated_path)
    return processed, found


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog=PROG_NAME,
        description="Add 'Original Path' and 'Updated Path' columns to a job-inventory workbook.",
    )
    parser.add_argument("--input", required=True, type=Path, help="Input .xlsx workbook.")
    parser.add_argument("--output", required=True, type=Path, help="Output .xlsx workbook.")
    parser.add_argument(
        "--updated-path",
        help="If given, written to the 'Updated Path' column of every row.",
    )
    parser.add_argument(
        "--log-level",
        default="INFO",
        choices=["DEBUG", "INFO", "WARNING", "ERROR"],
        help="Logging verbosity (default: INFO).",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    logging.basicConfig(
        level=args.log_level,
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
        filename=LOG_FILE,
        filemode="a",
    )

    with warnings.catch_warnings():
        # Some exporters omit the default cell style; openpyxl's fallback is harmless here.
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        wb = load_workbook(args.input)
    processed, found = update_workbook(wb.active, args.updated_path)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)

    logger.info("Rows processed: %d, original path found: %d", processed, found)
    logger.info("Workbook written to %s", args.output)
    print(f"Rows processed:       {processed}")
    print(f"Original path found:  {found}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
