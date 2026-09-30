"""CLI: annotate a Control-M job inventory workbook with command-path changes.

Reads a job inventory .xlsx (with a "Command Line" column) and the same
commands CSV used by git-pipeline-migration, then adds/updates two columns:
"Original Path" (the existing Command Line value) and "Updated Path" (the
replacement from the CSV, blank if that command has no update).
"""

from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path

import openpyxl

from git_pipeline_migration.commands import load_command_replacements

PROG_NAME = "update-job-listing"
LOG_FILE = f"{PROG_NAME}.log"

COMMAND_LINE_COLUMN = "Command Line"
ORIGINAL_PATH_COLUMN = "Original Path"
UPDATED_PATH_COLUMN = "Updated Path"

logger = logging.getLogger(__name__)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog=PROG_NAME,
        description=(
            "Add 'Original Path'/'Updated Path' columns to a job inventory "
            "workbook, populated from a commands CSV mapping."
        ),
    )
    parser.add_argument(
        "--xlsx-path",
        required=True,
        type=Path,
        help="Path to the job inventory .xlsx file (updated in place).",
    )
    parser.add_argument(
        "--commands-csv",
        required=True,
        type=Path,
        help="CSV file with 'Existing Command Line' and 'New Command Line Path' columns.",
    )
    parser.add_argument(
        "--log-level",
        default="INFO",
        choices=["DEBUG", "INFO", "WARNING", "ERROR"],
        help="Logging verbosity (default: INFO).",
    )
    return parser


def _header_index_map(worksheet) -> dict[str, int]:
    return {cell.value: cell.column for cell in worksheet[1] if cell.value is not None}


def _column_index(headers: dict[str, int], worksheet, name: str) -> int:
    if name in headers:
        return headers[name]
    column = worksheet.max_column + 1
    worksheet.cell(row=1, column=column, value=name)
    headers[name] = column
    return column


def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    logging.basicConfig(
        level=args.log_level,
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
        filename=LOG_FILE,
        filemode="a",
    )

    replacements = load_command_replacements(args.commands_csv)

    workbook = openpyxl.load_workbook(args.xlsx_path)
    worksheet = workbook.active

    headers = _header_index_map(worksheet)
    if COMMAND_LINE_COLUMN not in headers:
        raise ValueError(f"{args.xlsx_path} has no {COMMAND_LINE_COLUMN!r} column")
    command_col = headers[COMMAND_LINE_COLUMN]
    original_col = _column_index(headers, worksheet, ORIGINAL_PATH_COLUMN)
    updated_col = _column_index(headers, worksheet, UPDATED_PATH_COLUMN)

    rows_processed = 0
    rows_updated = 0
    for row in range(2, worksheet.max_row + 1):
        command_line = worksheet.cell(row=row, column=command_col).value or ""
        if not command_line and all(
            worksheet.cell(row=row, column=c).value is None
            for c in range(1, worksheet.max_column + 1)
        ):
            continue

        rows_processed += 1
        new_path = replacements.get(command_line, "")
        if new_path:
            rows_updated += 1
        worksheet.cell(row=row, column=original_col, value=command_line)
        worksheet.cell(row=row, column=updated_col, value=new_path)

    workbook.save(args.xlsx_path)

    logger.info("Rows processed: %d", rows_processed)
    logger.info("Rows with an updated path: %d", rows_updated)
    logger.info("Workbook written to %s", args.xlsx_path)

    print(f"Rows processed:           {rows_processed}")
    print(f"Rows with updated path:   {rows_updated}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
